#!/usr/bin/env python3
"""Build a Chinese TaskDec briefing from existing results; no GPU/inference."""
import csv
import io
import json
import re
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "analysis_exports/taskdec_submission_briefing_260911"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(exist_ok=True)
FONT = "Noto Sans CJK SC"
NAVY, INK, TEAL, MUTED = "142A3A", "213547", "00877F", "657785"
BG, WHITE, LINE, PALE = "F4F7F9", "FFFFFF", "DCE5EB", "E6F4F0"
ORANGE, RED, BLUE = "C47724", "B94D52", "466EB5"
W, H = 13.333333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "TaskDec｜现有结果与投稿讨论"
prs.core_properties.subject = "K-Radar、组件分析、VoD 适配与论文组织"
prs.core_properties.author = "TaskDec research"
prs.core_properties.keywords = "TaskDec, ICLR 2027, K-Radar, VoD"
SLIDES = []
SOURCES = {
    "main": "results/paper_main_table_kradar_v1_weather_compact_260901.md",
    "abl": "results/taskdec_v1_component_ablation_conf0_3_260901.md",
    "avail": "results/paper_availability_missing_modalities_260902.md",
    "vod": "results/paper_vod_main_table_draft_260909.md",
    "v2": "results/paper_supp_table_kradar_v2_generalization_260901.md",
    "eff": "results/paper_efficiency_table_260902.md",
    "review": "results/taskdec_project_and_paper_review_260909.md",
    "protocol": "results/taskdec_main_protocol_audit_vs_kradar_official_260910.md",
    "l4dr": "results/l4dr_vod_local_repro_results_260910.md",
    "l4drkr": "results/l4dr_reproduction_status_260910.md",
    "layout": "results/taskdec_figure_table_inventory_and_placement_260910.md",
    "gate": "analysis_exports/taskdec_bev_gate_fig4_260910/README.md",
    "pca": "analysis_exports/taskdec_patch_pca_weather_260909/metrics.json",
    "strength": "results/taskdec_v1_control_strength_compare_260904.md",
    "localasf": "results/exp_260818_202519_ASF_v1_0_local_repro_model2_full/summary_conf0.3.md",
    "conf0": "results/exp_260812_232650_TaskDecControlRobust_v1_model0_full/summary_conf0.0.md",
}


def tables(path):
    groups, block = [], []
    for line in (ROOT / path).read_text().splitlines() + [""]:
        if line.startswith("|"):
            cells = [re.sub(r"[*`]", "", c.strip()) for c in line.strip("|").split("|")]
            if not all(re.fullmatch(r"[:\- ]+", c) for c in cells):
                block.append(cells)
        elif block:
            groups.append(block)
            block = []
    return groups


def number(value):
    return float(value) if value not in ["-", ""] else None


main_rows = tables(SOURCES["main"])[0][1:]
weather03 = tables(SOURCES["main"])[1][1:]
weather05 = tables(SOURCES["main"])[2][1:]
main = {r[0]: r for r in main_rows}
our = main["TaskDec Robust (ours)"]
asf = main["ASF"]
weather_ours03 = next(r for r in weather03 if r[0] == "TaskDec Robust (ours)")
weather_asf03 = next(r for r in weather03 if r[0] == "ASF")
weather_ours05 = next(r for r in weather05 if r[0] == "TaskDec Robust (ours)")
weather_asf05 = next(r for r in weather05 if r[0] == "ASF")
weather_names = ["Normal", "Overcast", "Fog", "Rain", "Sleet", "Light snow", "Heavy snow"]
delta03 = [round(float(a)-float(b), 2) for a,b in zip(weather_ours03[3:], weather_asf03[3:])]
delta05 = [round(float(a)-float(b), 2) for a,b in zip(weather_ours05[3:], weather_asf05[3:])]
ablation = tables(SOURCES["abl"])[0][1:]
availability = tables(SOURCES["avail"])[0][1:]
efficiency = tables(SOURCES["eff"])[0][1:]
v2rows = tables(SOURCES["v2"])[0][1:]
metrics = json.loads((ROOT / SOURCES["pca"]).read_text())
assert round(float(our[6])-float(asf[6]), 2) == 8.05
assert round(float(our[4])-float(asf[4]), 2) == .31
assert delta03 == [8.09, .50, -.10, 7.93, .22, 8.39, -.30]
assert delta05 == [1.42, .55, .14, 7.59, -9.68, .70, -1.10]


def color(value):
    return RGBColor.from_string(value)


def box(s, x, y, w, h, fill=WHITE, line=None, radius=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        sh.adjustments[0] = .13
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(fill)
    if line:
        sh.line.color.rgb = color(line)
        sh.line.width = Pt(.8)
    else:
        sh.line.fill.background()
    return sh


def font_run(run, size, fill=INK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    prop = run._r.get_or_add_rPr()
    for tag in ["a:ea", "a:cs"]:
        item = OxmlElement(tag)
        item.set("typeface", FONT)
        prop.append(item)


def text(s, x, y, w, h, content, size=20, fill=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    for i, line in enumerate(str(content).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(5)
        p.line_spacing = 1.13
        run = p.add_run()
        run.text = line
        font_run(run, size, fill, bold)
    return sh


def picture(s, path, x, y, w, h, crop=False):
    # Fit within a slide viewport; no manipulation of scientific values/images.
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(w/iw, h/ih)
    dw, dh = iw*ratio, ih*ratio
    return s.shapes.add_picture(str(path), Inches(x+(w-dw)/2), Inches(y+(h-dh)/2),
                                width=Inches(dw), height=Inches(dh))


def arrow(s, points, fill=MUTED, width=1.6, dashed=False):
    for i, (a, b) in enumerate(zip(points, points[1:])):
        sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]),
                                    Inches(b[0]), Inches(b[1]))
        sh.line.color.rgb = color(fill)
        sh.line.width = Pt(width)
        ln = sh._element.spPr.get_or_add_ln()
        if dashed:
            dash = OxmlElement("a:prstDash")
            dash.set("val", "dash")
            ln.append(dash)
        if i == len(points)-2:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            ln.append(end)


def card(s, x, y, w, h, title, body, accent=TEAL, body_size=19):
    box(s,x,y,w,h,WHITE,LINE,True)
    box(s,x,y,.045,h,accent)
    short=h<1.5
    text(s,x+.22,y+(.10 if short else .17),w-.44,.5,title,18 if short else 22,accent,True)
    text(s,x+.22,y+(.57 if short else .85),w-.44,h-(.63 if short else .98),body,body_size)


def pill(s, x, y, w, label, fill=PALE, ink=TEAL):
    box(s,x,y,w,.34,fill,radius=True)
    text(s,x+.04,y+.015,w-.08,.28,label,11,ink,True,PP_ALIGN.CENTER)


def table(s, x, y, widths, headers, rows, row_h=.53, font=17, highlight=(),
          first_left=True, color_columns=None):
    sh = s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),
                            Inches(sum(widths)),Inches(row_h*(len(rows)+1)))
    tb = sh.table
    for col,w in zip(tb.columns,widths):
        col.width = Inches(w)
    for i,row in enumerate([headers]+rows):
        tb.rows[i].height=Inches(row_h)
        for j,value in enumerate(row):
            c=tb.cell(i,j)
            c.text=""
            c.fill.solid()
            c.fill.fore_color.rgb=color(NAVY if i==0 else PALE if i-1 in highlight else WHITE if i%2 else "EDF2F5")
            c.margin_left=c.margin_right=Inches(.10)
            c.margin_top=c.margin_bottom=Inches(.025)
            c.vertical_anchor=MSO_ANCHOR.MIDDLE
            p=c.text_frame.paragraphs[0]
            p.alignment=PP_ALIGN.LEFT if j==0 and first_left else PP_ALIGN.CENTER
            p.space_after=Pt(0)
            p.line_spacing=1.03
            run=p.add_run();run.text=str(value)
            ink=WHITE if i==0 else TEAL if i-1 in highlight else INK
            if i>0 and color_columns and j in color_columns:
                try:ink=TEAL if float(value)>0 else RED if float(value)<0 else MUTED
                except ValueError:pass
            font_run(run,font-1 if i==0 else font,ink,i==0 or i-1 in highlight)
    return sh


def slide(title, section, subtitle="", source="", notes="", refs=(), dark=False, compact=False):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=color(NAVY if dark else BG)
    i=len(prs.slides)
    if not dark:
        if compact:
            text(s,.55,.16,11.9,.4,title,20,INK,True)
        else:
            text(s,.55,.21,10.8,.3,section.upper(),11,TEAL,True)
            text(s,.55,.75,12.2,.63,title,29,INK,True)
            if subtitle:text(s,.57,1.45,12.1,.48,subtitle,16,MUTED)
        box(s,.55,7.09,12.22,.009,LINE)
        text(s,.57,7.17,11.55,.2,source,9.5,MUTED)
        text(s,12.22,7.14,.54,.25,f"{i:02d}",11,TEAL,True,PP_ALIGN.RIGHT)
    resolved=[SOURCES.get(r,r) for r in refs]
    body=notes+"\n\n资料来源：\n"+"\n".join(str(ROOT/r) for r in resolved)
    s.notes_slide.notes_text_frame.text=body
    SLIDES.append(dict(page=i,title=title,section=section,notes=notes,sources=resolved))
    return s


def takeaway(s, content, y=6.35, fill=PALE, ink=TEAL, size=18):
    box(s,.57,y,12.2,.50,fill,radius=True)
    text(s,.77,y+.09,11.8,.34,content,size,ink,True)


def metric_card(s,x,y,w,value,label,detail="",accent=TEAL):
    box(s,x,y,w,1.8,WHITE,LINE,True)
    text(s,x+.22,y+.12,w-.44,.66,value,37,accent,True)
    text(s,x+.22,y+.91,w-.44,.39,label,18,INK,True)
    if detail:text(s,x+.22,y+1.36,w-.44,.3,detail,13,MUTED)


def draw_delta_chart(s,x,y,w,h,title,values):
    text(s,x,y,w,.4,title,22,INK,True)
    label_w=1.48
    left=x+label_w;right=x+w-.40;zero=(left+right)/2;scale=(right-left)/20
    top=y+.70;step=(h-1.2)/7
    for tick in [-10,-5,0,5,10]:
        xx=zero+tick*scale
        box(s,xx,top-.10,.009,step*7+.08,"BDCCD4" if tick==0 else LINE)
        text(s,xx-.25,top+step*7+.09,.5,.25,str(tick),11,MUTED,align=PP_ALIGN.CENTER)
    for i,(name,v) in enumerate(zip(weather_names,values)):
        yy=top+i*step
        text(s,x,yy-.015,label_w-.1,.3,name,14,MUTED)
        bx=zero if v>=0 else zero+v*scale
        box(s,bx,yy+.035,max(abs(v)*scale,.015),.22,TEAL if v>=0 else RED)
        if v>=0:tx=zero+v*scale+.06
        else:tx=zero+v*scale-.70
        text(s,tx,yy-.025,.69,.34,f"{v:+.2f}",13,TEAL if v>=0 else RED,True)


def make_pca():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rows=[]
    source=ROOT/"analysis_exports/taskdec_patch_pca_weather_260909/taskdec_patch_states_pca.csv"
    with source.open() as f:
        for r in csv.DictReader(f):
            if r["foreground"]=="1":rows.append(r)
    features=["raw","common","unique"]
    colors={"Camera":"#C5524F","LiDAR":"#248A76","4D Radar":"#466EB5"}
    rng=np.random.default_rng(2026)
    fig,axes=plt.subplots(2,3,figsize=(10.6,4.7))
    shown={}
    for j,feat in enumerate(features):
        allf=[r for r in rows if r["feature_type"]==feat]
        allxy=np.array([[float(r['pc1']),float(r['pc2'])] for r in allf])
        lo,hi=allxy.min(axis=0),allxy.max(axis=0);pad=(hi-lo)*.08
        for i,weather in enumerate(["normal","rain"]):
            ax=axes[i,j]
            for mod,c in colors.items():
                cand=[r for r in allf if r['weather']==weather and r['modality']==mod]
                idx=rng.choice(len(cand),min(len(cand),190),replace=False)
                pts=np.array([[float(cand[k]['pc1']),float(cand[k]['pc2'])] for k in idx])
                ax.scatter(pts[:,0],pts[:,1],s=8,c=c,alpha=.68,lw=0,label=mod)
                shown[f'{feat}/{weather}/{mod}']=len(idx)
            ax.set_xlim(lo[0]-pad[0],hi[0]+pad[0]);ax.set_ylim(lo[1]-pad[1],hi[1]+pad[1])
            ax.set_xticks([]);ax.set_yticks([])
            if i==0:ax.set_title(['Canonical token','Common state','Unique state'][j],fontsize=14,pad=10)
            if j==0:ax.set_ylabel(weather.title(),fontsize=13,labelpad=13)
            for spine in ax.spines.values():spine.set_color('#DCE5EB')
    handles,labels=axes[0,0].get_legend_handles_labels()
    fig.legend(handles,labels,loc='lower center',ncol=3,frameon=False,fontsize=12)
    fig.subplots_adjust(left=.08,right=.995,top=.91,bottom=.10,hspace=.17,wspace=.12)
    fig.savefig(ASSETS/'pca_normal_rain.png',dpi=180,facecolor='white')
    plt.close(fig)
    return dict(source=str(source),displayed_points=shown,
                projection='Existing PCA coordinates, separately fitted by representation in original exporter.',
                limits='All-foreground min/max per representation, same limits across weather; no outlier clipping.',
                seed=2026)


PCA_INFO=make_pca()
GATE_DIR=ROOT/'analysis_exports/taskdec_bev_gate_fig4_260910'
VIS_DIR=ROOT/'analysis_exports/taskdec_paper_visuals_260909'
# Camera context is embedded directly from the original data; no raw-file copy.
manifest=json.loads((GATE_DIR/'manifest.json').read_text())
rain_frame=next(f for f in manifest['frames'] if f['id']=='seq25_rdr00154')
with Image.open(rain_frame['camera_path']) as im:
    buf=io.BytesIO();im.crop((0,0,1280,720)).convert('RGB').save(buf,format='JPEG',quality=89);buf.seek(0)
CAMERA_BYTES=buf.getvalue()


def camera(s,x,y,w,h):
    # Fixed 16:9 front0 view; no photometric enhancement.
    stream=io.BytesIO(CAMERA_BYTES)
    sh=s.shapes.add_picture(stream,Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    return sh


def build():
    s=slide('TaskDec｜现有结果与投稿讨论','开场',dark=True,
            notes='建议用约 15–20 分钟讲完前 18 页；后 6 页按问题翻阅。核心结果来自既有实验，本次只整理汇报，不新增训练或推理。主对照是官方 ASF checkpoint 的同置信度复评。',refs=['main','vod','gate'])
    text(s,.62,.45,10,.4,'RESEARCH BRIEFING  /  ICLR 2027 投稿讨论稿',14,'94C9C6',True)
    text(s,.60,1.36,6.6,1.0,'TaskDec',59,WHITE,True)
    text(s,.65,2.6,6.7,1.55,'任务感知解耦控制\n多传感器 3D 目标检测',30,WHITE,True)
    text(s,.67,4.12,6.4,.75,'方法贡献 · 实验依据 · 论文安排',20,'A8BFCD')
    camera(s,7.45,1.5,5.26,2.959)
    text(s,7.5,4.6,5.1,.5,'K-Radar 真实雨夜场景 / seq25',13,'A8BFCD')
    for x,val,lab in [(0.65,'88.36','K-Radar AP3D@0.3'),(4.91,'+8.05','相对官方 ASF / conf=0.3'),(9.17,'+0.30','VoD EAA / 相对 PP-Concat')]:
        box(s,x,5.38,3.53,1.34,'203E50',radius=True)
        text(s,x+.2,5.49,3.15,.57,val,34,WHITE,True)
        text(s,x+.2,6.18,3.15,.32,lab,13,'B7CED8')
    text(s,.67,7.06,10,.23,'资料整理：2026.09.11  ｜  已纳入 9 月 10 日 L4DR 复现与协议审计',11,'9CB5C4')

    s=slide('现有证据可以形成一条完整论文主线','01 / 结果概览',
            '从“表征分离”到“局部融合控制”，再到主基准与骨干适配。',
            '依据：主表、组件消融、VoD 官方指标汇总',
            '先讲三句话：主基准在已收集比较中有竞争力；四项移除都使 AP3D@0.3 下降；VoD 支持架构适配和相对原生强基线的 EAA 小幅收益。不要把三条证据夸大成各天气、各阈值、各数据集全面最优。',
            ['main','abl','vod','l4dr'])
    card(s,.57,2.12,3.9,3.6,'主基准有清楚收益','K-Radar v1.0\n88.36 AP3D@0.3\n官方 ASF：80.31\n\nAP3D@0.5：+0.31',body_size=22)
    card(s,4.72,2.12,3.9,3.6,'组件确实参与工作','移除四类组件后\nAP3D@0.3 下降\n8.13–8.74 点\n\n已有小测选择＋全量结果',body_size=22)
    card(s,8.87,2.12,3.9,3.6,'可适配到 VoD','原生 PP 风格骨干\nEAA：69.88 → 70.18\nDC：83.80 → 83.79\n\nL4DR 绝对性能仍更高',body_size=22)
    takeaway(s,'讨论重点：贡献定位是否充分，以及最值得补的证据是什么。')

    s=slide('动机：空间对齐之后，局部信息仍需任务控制','02 / 研究问题',
            '同一 patch 内，不同传感器的观测质量与任务相关性可能不同。',
            '依据：方法梳理；相机为 K-Radar seq25/rdr00154，未做图像增强',
            '这页是汇报用动机示意，不是已经完成的论文 Fig.1。照片只展示真实场景中的雨滴和灯光干扰，不能据此判断某一传感器的真实可靠性。研究问题是：对齐后，哪些状态参与控制，控制发生在什么位置。',
            ['review','gate'])
    camera(s,.58,2.15,6.65,3.741)
    pill(s,.77,2.34,2.6,'真实场景参照 / 雨夜',NAVY,WHITE)
    card(s,7.57,2.15,5.18,1.15,'已解决的基础','各传感器映射到统一 patch 空间',body_size=17)
    card(s,7.57,3.56,5.18,2.32,'TaskDec 要回答','哪些表征用于传感器缩放？\n哪些位置需要更强控制？\n任务信息如何进入 attention？',body_size=21)
    takeaway(s,'将 common / unique 状态用于前景门控、模态缩放和任务上下文调制。')

    s=slide('TaskDec 的新增部分作用于融合输入、query 和输出','03 / 方法框架',
            '以 ASF canonical patch fusion 为基础，学习用于控制融合的 common / unique 状态。',
            '依据：patch_dec_a2_fusion.py；fusion_base_integrated.py；主配置',
            '这是按当前实现重画的可编辑汇报示意。K/V 来自经过 gate 和 sensor scale 控制的原始 token 与 dec residual；context 在 attention 前调制 query，并在 attention 后、PFT 前注入残差。GT 仅用于训练监督。详式见下一页。',
            ['review','models/fuser/patch_dec_a2_fusion.py','models/skeletons/fusion_base_integrated.py'])
    for y,label,c in [(2.65,'Camera',BLUE),(3.51,'LiDAR',TEAL),(4.37,'4D Radar',ORANGE)]:
        box(s,.58,y,1.40,.62,c,radius=True);text(s,.62,y+.15,1.32,.28,label,16,WHITE,True,PP_ALIGN.CENTER)
    box(s,2.28,3.08,1.72,1.48,WHITE,LINE,True)
    text(s,2.43,3.28,1.42,1.1,'Encoder +\ncanonical\npatch token x',17,INK,True,PP_ALIGN.CENTER)
    for yy in [2.96,3.82,4.68]:arrow(s,[(1.98,yy),(2.12,yy),(2.12,3.82),(2.28,3.82)])
    box(s,4.32,3.08,1.69,1.48,WHITE,LINE,True)
    text(s,4.42,3.33,1.49,1.0,'Common c\nUnique u\n每模态独立 MLP',17,TEAL,True,PP_ALIGN.CENTER)
    arrow(s,[(4.0,3.82),(4.32,3.82)])
    box(s,6.31,3.08,1.98,1.48,PALE,TEAL,True)
    text(s,6.43,3.32,1.74,1.03,'解耦状态控制器\ngate / reliability\ntask context',17,TEAL,True,PP_ALIGN.CENTER)
    arrow(s,[(6.01,3.82),(6.31,3.82)],TEAL)
    for y,lab in [(2.25,'受控 K/V：x′'),(3.62,'调制 query：q′'),(5.0,'输出 context 残差')]:
        box(s,8.67,y,1.92,.7,WHITE,LINE,True)
        text(s,8.75,y+.17,1.76,.37,lab,15,INK,True,PP_ALIGN.CENTER)
        arrow(s,[(8.29,3.82),(8.48,3.82),(8.48,y+.35),(8.67,y+.35)],TEAL)
    box(s,11.01,3.08,1.74,1.18,NAVY,radius=True)
    text(s,11.1,3.31,1.56,.8,'Patch\nattention',20,WHITE,True,PP_ALIGN.CENTER)
    arrow(s,[(10.59,2.60),(11.87,2.60),(11.87,3.08)])
    arrow(s,[(10.59,3.97),(11.01,3.97)])
    box(s,11.01,4.74,1.74,.75,PALE,TEAL,True)
    text(s,11.08,4.94,1.6,.3,'+ context 残差',16,TEAL,True,PP_ALIGN.CENTER)
    arrow(s,[(11.87,4.26),(11.87,4.74)])
    arrow(s,[(10.59,5.35),(10.79,5.35),(10.79,5.12),(11.01,5.12)],TEAL)
    arrow(s,[(11.87,5.49),(11.87,5.77)])
    box(s,10.78,5.77,1.97,.50,NAVY,radius=True)
    text(s,10.84,5.88,1.85,.29,'PFT → 检测头',15,WHITE,True,PP_ALIGN.CENTER)
    text(s,2.34,5.00,5.9,.92,'x 与分解状态共同构造受控 K/V；\n监督项不位于推理前向路径上。',18,MUTED)
    takeaway(s,'训练监督：common 对齐、unique 分离、分支去相关、前景与任务目标。',y=6.46,size=17)

    s=slide('三路控制各有落点，gate 与 context 的语义要区分','04 / 关键计算',
            '下面保留控制关系；正式公式、超参数和损失权重放入方法与附录。',
            '依据：主方法代码、主配置；v1.0 context 为单类 objectness context',
            'g 决定控制强弱，r 在可用模态之间归一化，h 提供任务方向。v1 主实验只启用 Sedan，因此 context 是 binary objectness 的投影。common/unique 是学习状态，不保证 x=c+u，也不能把 unique 直接叫噪声。均匀 reliability 只令 scale 为 1，残差和 context 仍存在。',
            ['review','configs/ASF_task_dec_controlled_robust_v1_0.yml'])
    card(s,.57,2.1,3.9,3.62,'01  前景 gate','由 common 均值与\n|unique| 均值预测 g\n\n回答：这个 patch\n需要多强的控制？',body_size=21)
    card(s,4.72,2.1,3.9,3.62,'02  传感器缩放','s = clip[1 + γg(Mr − 1)]\nx′ = s[x + αg(c + u)]\n\n控制 attention 的 K/V\n保留原始 token 的路径',body_size=19)
    card(s,8.87,2.1,3.9,3.62,'03  任务 context','q′ = q + βq · g · h\nz′ = Attn(q′, x′, x′)\n        + βz · g · h\n\n注入 query 与融合输出',body_size=19)
    takeaway(s,'v1.0 是 objectness context；reliability 是学习权重，目前没有真实可靠性标签。',size=17)

    s=slide('先锁定主实验的比较口径','05 / 评测设置',
            '主结果统一对应 K-Radar v1.0：窄 ROI、Sedan、官方 train/test 划分。',
            '依据：9 月 10 日主协议审计；数值来源与置信度在每张比较表注明',
            '应区分 IoU 阈值和置信度过滤阈值。主对照统一 conf_thr=0.3；ASF 用发布 checkpoint 本地复评；其他方法来自已经整理的公开论文或日志。L4DR 本地 K-Radar 强制适配输入的结果是协议不匹配诊断，不用来排名。',
            ['protocol','main','l4drkr'])
    table(s,.57,2.10,[2.35,5.77,4.08],['项目','主设定','汇报中如何标注'],[
        ['任务 / 标签','Sedan / label v1.0','单类主基准'],
        ['空间范围','x: 0–72 m；y: ±6.4 m','窄行驶走廊 ROI'],
        ['输入传感器','Camera + LiDAR + 4D Radar','C + L + R'],
        ['主模型','TaskDec Robust / model_0','主方法完整配置'],
        ['报告指标','AP3D / APBEV，IoU 0.3 / 0.5','conf_thr = 0.3'],
        ['ASF 对照','官方发布 checkpoint 本地复评','与 TaskDec 同置信度过滤'],
    ],row_h=.53,font=18)
    takeaway(s,'其他文献行保留来源；不把公开结果表述成全部方法的同环境重跑。',size=18)

    s=slide('K-Radar：AP3D@0.3 达到 88.36','06 / 主结果',
            '与官方 ASF 同阈值比较：AP3D@0.3 +8.05，AP3D@0.5 +0.31。',
            '依据：v1 compact 主表；ASF / TaskDec conf=0.3；其余为已收集文献数值',
            '这一页是主要性能证据。直接基线是官方 ASF checkpoint。L4DR 是相关 L+R 专用方法，但模态集合不同；AW-MoE-LRC 提供另一项 C+L+R 的公开参考。在已收集的表中 TaskDec 总体 AP3D@0.3 最强，不能延伸为所有任务、阈值全面领先。',
            ['main','protocol'])
    rows=[]
    for r in main_rows:
        name='TaskDec' if r[0].startswith('TaskDec') else r[0]
        src='本地 / 主模型' if name=='TaskDec' else '官方 ckpt 复评' if name=='ASF' else '文献'
        rows.append([name,r[1],src,r[6],r[4]])
    table(s,.57,2.10,[3.02,1.58,2.74,2.43,2.43],
          ['方法','模态','来源','AP3D@0.3','AP3D@0.5'],rows,row_h=.47,font=18,highlight=[6])
    takeaway(s,'论文主张聚焦：指定 v1.0 协议下的总体表现与任务控制收益。',y=6.40)

    s=slide('天气收益主要来自 Normal、Rain 与 Light snow','07 / 天气分解',
            '以下都是 TaskDec − 官方 ASF 的 AP3D 绝对差值，单位为百分点。',
            '依据：v1 weather compact 表；精确数值见备份页 20',
            '不要说所有恶劣天气都提升。IoU=0.3 下 Normal、Rain、Light snow 增益突出，Fog 与 Heavy snow 略低。IoU=0.5 下 Rain 仍有收益，但 Sleet 下降 9.68 点。全量 AP 不是七种天气 AP 的简单平均。',
            ['main'])
    draw_delta_chart(s,.58,2.03,6.0,4.1,'IoU = 0.3',delta03)
    draw_delta_chart(s,6.85,2.03,5.90,4.1,'IoU = 0.5',delta05)
    takeaway(s,'严格 IoU 下的 Sleet 是需要交代的局限，不能概括成“全天气一致提升”。',size=17)

    s=slide('四项移除均明显降低 AP3D@0.3','08 / 组件消融',
            '完整模型 88.36；移除后为 79.62–80.23，说明各部分都在当前配置下发挥作用。',
            '依据：组件消融结果；完整模型与移除项均经过 1000 样本小测选择',
            '作者已确认移除项也先做了 1000 样本小测，再选对应最优对比并全量评测，不能因为 model_0 与 model_9 编号不同就判定选择不一致。后续需固化样本清单、选择指标、候选范围。这里支持所报告配置下组件有效，不声称统计显著性或线性可加贡献。',
            ['abl'])
    labels=['完整 TaskDec','移除 sensor reliability','移除 task context','移除解耦监督','移除 foreground gate']
    rows=[[lab,r[1],r[8],r[7],r[9]] for lab,r in zip(labels,ablation)]
    table(s,.57,2.16,[4.35,1.60,2.08,2.08,2.09],
          ['变体','checkpoint','AP3D@0.3','AP3D@0.5','Δ @0.3'],rows,row_h=.60,font=18,highlight=[0],color_columns={4})
    takeaway(s,'已确认选择流程；投稿前补齐小测样本、选择指标及候选范围的复现记录。',size=17)

    s=slide('PCA 提供表征结构线索，不能代替功能证据','09 / 表征诊断',
            '展示 Normal / Rain 的前景 patch；common 的跨模态 cosine 也可作为高维补充。',
            '依据：168 帧原始 PCA 导出；每面板每模态最多 190 个前景点，种子固定',
            'PCA 使用已有投影坐标，不重新训练或拟合；按表征分别采用覆盖全部前景点的轴范围，同一表征两天气共用轴范围。二维图不能证明统计独立或物理语义。高维前景 common 跨模态 cosine 均值约 0.948–0.962。功能作用主要由消融和检测结果支持。',
            ['pca','analysis_exports/taskdec_patch_pca_weather_260909/taskdec_patch_states_pca.csv'])
    picture(s,ASSETS/'pca_normal_rain.png',.58,2.04,9.65,4.60)
    card(s,10.42,2.18,2.32,3.98,'诊断读法','结构可观察\n\ncommon\ncosine 均值\n0.948–0.962\n\n不能直接解释\n为“纯目标 / 噪声”',body_size=17)

    s=slide('空间 gate：真实场景与完整 patch 输出','10 / 控制可视化',
            source='依据：Fig.4 完整 gate 导出；16×90 / 帧；统一 [0,1]；GT 仅作参照',
            notes='这页直接使用昨天完成的 Fig.4。左侧阴天高速展示两处不同距离的车辆响应；右侧雨夜路口有多个近距响应，较远车辆的响应更弱。每帧完整 1440 个 gate，正文两帧独立复算逐项一致。只说明所选场景的空间行为，不是 ASF–TaskDec 检测对比。另有一帧小雪候选复算不稳定，未使用。',
            refs=['gate','analysis_exports/taskdec_bev_gate_fig4_260910/validation.json'],compact=True)
    picture(s,GATE_DIR/'paper_fig4_taskdec_bev_gate_draft.png',.26,.54,12.82,6.45)

    s=slide('缺模态：LR / LC 总体可用，RC 是主要失效组合','11 / 可用模态变化',
            '所有行使用 C+L+R 训练的 checkpoint，仅控制推理时可用的传感器集合。',
            '依据：missing-modality 表；LR 使用已保存预测的重算结果；conf=0.3',
            'LR 与 LC 的 AP3D@0.3 比官方 ASF 更高。LR 在 AP3D@0.5 略低 0.30，LC 在两个阈值都更高。RC 没有 LiDAR，两个阈值均明显退化。这和当前 reliability 读数偏 LiDAR 的观察相容，但不能只凭相关性断定因果。',
            ['avail','pca'])
    ar={(r[0].startswith('TaskDec'),r[1]):r for r in availability}
    rows=[]
    for mode in ['RLC','LR','LC','RC']:
        a,b=ar[(True,mode)],ar[(False,mode)]
        rows.append([mode,b[5],a[5],f'{float(a[5])-float(b[5]):+.2f}',
                     f'{float(a[3])-float(b[3]):+.2f}'])
    table(s,.57,2.25,[2.0,2.6,2.6,2.5,2.5],
          ['可用模态','ASF @0.3','TaskDec @0.3','Δ AP3D@0.3','Δ AP3D@0.5'],rows,
          row_h=.66,font=21,color_columns={3,4})
    text(s,.72,5.82,11.7,.38,'C：相机    L：LiDAR    R：4D Radar',17,MUTED)
    takeaway(s,'正文简述 LR / LC 的正向结果与 RC 例外，完整表保留在附录。',size=18)

    s=slide('VoD：适配原生 PP 骨干后，EAA 超过强拼接基线','12 / 跨数据集与骨干适配',
            '泛化含义：迁移架构并在 VoD 上训练；主指标使用官方 EAA / DC。',
            '依据：VoD 主表；已纳入 9 月 10 日 L4DR 本地复现，详见备份页 22',
            '按既定写作重心，先讲方法可用于原生 PointPillars 风格 LiDAR–4D radar 特征，再讲相对强 PP-Concat 的 EAA +0.30。DC 基本持平。当前是 mild + warm start、单次选定运行；PP-Concat 用 AMP，TaskDec 用 FP32，训练差异要披露。L4DR 文献和本地复现绝对性能均更高，一句话说明并把表放备份。',
            ['vod','l4dr'])
    card(s,.57,2.15,4.2,2.50,'从 K-Radar 到 VoD','K-Radar：C + L + R\ncanonical patch fusion\n\nVoD：L + 4DR\n原生 PP 风格融合',body_size=20)
    table(s,5.13,2.17,[3.68,1.94,1.94],['本地方法','EAA mAP','DC mAP'],[
        ['PP-Concat / ep80','69.88','83.80'],
        ['TaskDec-PP / ep79','70.18','83.79'],
    ],row_h=.73,font=20,highlight=[1])
    metric_card(s,5.13,4.70,3.67,'+0.30','EAA mAP','强 PP-Concat 上的适配收益')
    metric_card(s,9.03,4.70,3.66,'−0.01','DC mAP','保持基本相当',accent=MUTED)
    text(s,.74,5.04,3.81,1.44,'当前配置\nmild + PP-Concat warm start\nTaskDec FP32 / 基线 AMP',16,MUTED)
    text(s,.73,6.67,11.8,.23,'L4DR 仍更强：文献 72.70 / 87.47；本地复现 71.00 / 84.84（EAA / DC）。',13,MUTED)

    s=slide('v2 补充结果来自早期 DecControlled 变体','13 / 扩展材料，拟入论文附录',
            '宽 ROI、Sedan + Bus/Truck、C+L+R；该变体没有完整 TaskDec 的 task-context 分支。',
            '依据：v2 supplementary 表；conf=0.3；所选天气均值不是 Total',
            '这组结果明确标注为 early variant，不能说完整 v1 TaskDec 不变迁移。总量 AP3D@0.5 对 Sedan 近似持平，对 Bus/Truck 有收益；selected-weather 有更清楚的正向结果，但它只平均四种天气。AP3D@0.3 的 Sedan Total 稍低。论文放附录，正文一两句说明。',
            ['v2'])
    pill(s,.59,2.05,4.4,'模型身份：DecControlled / early variant','F9ECD9',ORANGE)
    table(s,.57,2.69,[2.15,2.5,2.5,2.5,2.55],
          ['类别','ASF Total','早期变体 Total','ASF 所选天气','早期变体 所选天气'],[
              ['Sedan',v2rows[0][2],v2rows[2][2],v2rows[0][7],v2rows[2][7]],
              ['Bus / Truck',v2rows[1][2],v2rows[3][2],v2rows[1][7],v2rows[3][7]],
          ],row_h=.75,font=18)
    text(s,.76,5.18,11.72,.90,'表内指标：AP3D@0.5\n所选天气均值：Overcast、Rain、Light snow、Heavy snow，排除 Total。',19,MUTED)
    takeaway(s,'定位为方法家族的扩展证据；完整模型的跨骨干适配由 VoD 实验承担。',size=18)

    s=slide('新增控制带来约 1.7% 参数与 8.1% 前向耗时','14 / 推理开销',
            '同机、batch size 1；20 次 warmup，100 次计时。',
            '依据：efficiency 汇总及对应 JSON；前向耗时与端到端耗时分开报告',
            '参数增量 1.34M；forward +6.50 ms。只比较 ASF 与完整 TaskDec 的相同 C+L+R 设置。端到端 647.50→687.56 ms，显著受数据加载影响；forward FPS 不能叫端到端 FPS。这里不作实时系统承诺。',
            ['eff','results/efficiency_260902/taskdec_robust_v1_model0_rlc.json'])
    rows=[[r[0].split(' v1')[0].replace('TaskDec Robust model_0 (C+L+R)','TaskDec'),r[1],r[2],r[3],r[5]] for r in efficiency[:2]]
    rows[0][0]='官方 ASF';rows[1][0]='TaskDec'
    table(s,.57,2.16,[2.8,2.35,2.4,2.3,2.35],['模型','参数 (M)','Forward (ms)','Forward FPS','峰值显存 (GB)'],rows,
          row_h=.74,font=20,highlight=[1])
    metric_card(s,.59,4.74,3.83,'+1.34 M','参数增量','78.13 → 79.47 M')
    metric_card(s,4.75,4.74,3.83,'+6.50 ms','前向耗时增量','79.95 → 86.45 ms')
    metric_card(s,8.91,4.74,3.83,'687.56 ms','TaskDec 端到端耗时','含数据加载；单独报告',accent=MUTED)

    s=slide('论文视觉主线按 5 张图、4 个紧凑表组织','15 / 正文与附录安排',
            '现有数字已较完整，接下来主要是图稿、论证顺序与版面取舍。',
            '依据：9 月 10 日图表清单；这是当前写作规划，尚未验证完整 LaTeX 占页',
            '保留独立动机、主框架、PCA、空间 gate 和实车检测对比。Fig.4 已有实际导出；Fig.2 有草稿但应纠正训练监督箭头。Fig.1 和 Fig.5 尚未完成论文稿。四个主表为 v1 主比较、天气、组件、VoD。中心距离和 reliability、缺模态全表、早期 v2 以及协议、效率进入附录。',
            ['layout','gate'])
    figs=[('Fig.1','独立动机','待制作',ORANGE),('Fig.2','主框架','草稿待修',ORANGE),
          ('Fig.3','代表 PCA','数据已有',TEAL),('Fig.4','空间 gate','初稿已有',TEAL),('Fig.5','实车检测对比','待制作',ORANGE)]
    for i,(num,name,status,c) in enumerate(figs):
        x=.57+i*2.48
        box(s,x,2.15,2.29,1.45,WHITE,LINE,True)
        text(s,x+.15,2.32,1.98,.33,num,18,TEAL,True)
        text(s,x+.15,2.85,1.98,.43,name,18,INK,True)
        pill(s,x,3.81,2.29,status,PALE if c==TEAL else 'F9ECD9',c)
    text(s,.63,4.53,2.16,.46,'正文四表',22,INK,True)
    for i,label in enumerate(['v1 主结果','天气分解','组件消融','VoD 官方指标']):
        pill(s,2.91+i*2.46,4.57,2.24,label,NAVY,WHITE)
    card(s,.59,5.33,12.14,1.14,'附录承接','完整表与协议、缺模态 / v2、全天气 PCA、中心距离 / reliability、效率和超参数。',body_size=17)

    s=slide('明天希望与师兄明确的四个问题','16 / 投稿讨论议程',
            '把讨论落到贡献判断、证据缺口与稿件范围。',
            '依据：现有方法、结果、图表与协议梳理；以下是待讨论事项',
            '建议让师兄先评价机制贡献是否足以支撑目标会议，再决定最值得补的实验，而不是泛泛追加所有实验。比较协议主口径已确定使用官方 ASF checkpoint；仍需把来源和敏感性说明准备好。VoD 的强项是可适配，收益幅度小，应确认这种证据在目标投稿中是否足够。',
            ['review','layout','protocol','l4dr'])
    card(s,.57,2.11,5.95,1.78,'1  贡献是否讲得成立？','“解耦状态 → 局部控制”能否成为核心贡献？\n与已有解耦、query 调制、可靠性融合如何区分？',body_size=19)
    card(s,6.81,2.11,5.95,1.78,'2  最值得补哪类证据？','组件重复运行，还是更直接的受控诊断？\n优先回答哪一个潜在审稿问题？',body_size=19)
    card(s,.57,4.34,5.95,1.78,'3  泛化材料是否足够？','VoD 的 +0.30 EAA 与 DC 持平如何定位？\nRC 退化和 v2 早期变体说明到什么程度？',body_size=19)
    card(s,6.81,4.34,5.95,1.78,'4  正文取舍与投稿节奏？','5 图 4 表如何压缩并保留真实场景展示？\n先定稿主线，再确定补实验清单与停止条件。',body_size=19)
    takeaway(s,'目标：讨论结束时明确“主贡献一句话、补证据优先级、正文范围”。',size=18)

    s=slide('下一步先完成可直接进入论文的材料','17 / 建议行动顺序',
            '以下是建议优先级，待明天讨论后确认投入。',
            '依据：现有材料完成度；未据此启动新训练或评测',
            '第一优先级是把已有内容变成可审的图表和文字：修框架图，补动机和同帧检测对比，统一表注和 checkpoint 身份。第二优先级固定选择流程、数据协议和代码快照。是否重复主模型与消融、扩展缺模态或强化 VoD，由讨论后的研究问题和资源决定，不把尚未运行的事情说成计划已获批准。',
            ['layout','abl','protocol','gate'])
    table(s,.57,2.18,[1.20,3.65,5.19,2.16],['优先级','任务','完成物','当前状态'],[
        ['P0','主框架与独立动机','准确区分前向计算、训练监督与研究动机','已有素材'],
        ['P0','同帧 ASF–TaskDec 检测对比','真实相机 / BEV、统一阈值、预测与 GT','待制作'],
        ['P0','正文与图注','主结果 → 组件 → VoD；Fig.4 可先入稿','可推进'],
        ['P1','复现与选择记录','小测样本 / 指标 / 候选范围 / 模型快照','待固化'],
        ['待讨论','补强关键证据','优先选择能改变审稿判断的一组实验','未启动'],
    ],row_h=.67,font=17)
    takeaway(s,'先用现有证据形成完整初稿，再围绕最关键的问题补强。',y=6.45)

    s=slide('备份 A｜主比较的完整四项 AP 指标','BACKUP / 完整主表',
            'AP 单位为 %；“文献”行是已收集的公开报告，ASF 为官方 checkpoint 同阈值复评。',
            '依据：v1 compact 主表；C=相机，L=LiDAR，R=4D Radar；conf=0.3',
            '这张表用于回答 BEV 指标与 IoU=0.5 的问题。L4DR 的公开数值来自文献汇总；本地强制 label v1.0 和替代 radar sparse 的跑法未协议对齐，所以不把 24.80 AP3D@0.3 的诊断结果纳入论文排名。',
            ['main','l4drkr','protocol'])
    rows=[]
    for r in main_rows:
        rows.append(['TaskDec' if r[0].startswith('TaskDec') else r[0],r[1],r[3],r[4],r[5],r[6]])
    table(s,.57,2.16,[3.05,1.38,1.94,1.94,1.94,1.95],
          ['方法','模态','BEV@0.5','3D@0.5','BEV@0.3','3D@0.3'],rows,row_h=.47,font=17,highlight=[6])
    takeaway(s,'L4DR 的本地 K-Radar 协议错配诊断不进入有效性能比较。',y=6.44,size=18)

    s=slide('备份 B｜天气分解的精确数值','BACKUP / 天气表',
            '同一官方 ASF checkpoint 与 TaskDec 主模型；保留两种 IoU 阈值。',
            '依据：v1 weather compact 表；所有差值由表内数值相减得到',
            '按读者提问查数。AP@0.3 的 Rain +7.93；AP@0.5 的 Rain +7.59。Sleet 的 IoU=0.3 基本相当，但 IoU=0.5 -9.68。不要把 Sleet 与 Rain / Light snow 的强收益混为一谈。',
            ['main'])
    rows=[]
    for i,name in enumerate(weather_names):
        rows.append([name,weather_asf03[i+3],weather_ours03[i+3],f'{delta03[i]:+.2f}',
                     weather_asf05[i+3],weather_ours05[i+3],f'{delta05[i]:+.2f}'])
    table(s,.57,2.18,[2.42,1.63,1.63,1.63,1.63,1.63,1.63],
          ['天气','ASF @0.3','Ours @0.3','Δ @0.3','ASF @0.5','Ours @0.5','Δ @0.5'],
          rows,row_h=.49,font=16,color_columns={3,6})
    takeaway(s,'Total 是全测试集 AP，不等于这里七行的算术平均。',y=6.48,size=18)

    s=slide('备份 C｜官方 checkpoint 与置信度敏感性','BACKUP / 比较口径',
            '主比较使用官方 ASF checkpoint；其他置信度和本地重训结果作为敏感性材料。',
            '依据：主协议审计、conf=0.0 全量汇总、本地 ASF 重训汇总',
            '主动准备这张表应对 baseline 选择问题。官方 headline 接近 conf=0.0 的结果，主表是同 conf=0.3 复评，不能拿 conf=0.3 的 +8.05 去宣称相对所有 ASF 跑法都有这个增益。本地 ASF 重训 AP3D@0.3 为 88.57，高于 TaskDec 的 88.36；因此保留训练与选择敏感性说明。它作为附录材料，不替换已决定使用的官方 checkpoint 主表。',
            ['protocol','conf0','localasf'])
    table(s,.57,2.13,[5.3,1.6,2.65,2.65],['模型 / 来源','conf_thr','AP3D@0.3','AP3D@0.5'],[
        ['ASF 官方日志 / exp250219','0.0','87.42','73.58'],
        ['ASF 官方日志 / exp250303','0.0','87.34','72.95'],
        ['TaskDec 主模型','0.0','88.06','72.83'],
        ['ASF 官方 ckpt 复评 / 主对照','0.3','80.31','67.19'],
        ['TaskDec 主模型 / 主表','0.3','88.36','67.50'],
        ['ASF 本地重训 / model_2','0.3','88.57','67.49'],
    ],row_h=.52,font=17,highlight=[3,4])
    takeaway(s,'限定比较对象与协议；不宣称对所有 ASF 训练结果或阈值都存在大幅优势。',size=17)

    s=slide('备份 D｜VoD 中的 L4DR 本地复现与训练差异','BACKUP / 9 月 10 日新增',
            '完整 L4DR 的文献结果和本地复现都高于当前 TaskDec-PP。',
            '依据：l4dr_vod_local_repro_results_260910.md；VoD 官方 EAA / DC',
            '新增 L4DR 运行正常完成 100 epochs，最终两次评测中 epoch99 更好：71.00 EAA /84.84 DC；不是证明对全部 100 个 epochs 搜索后的全局最优。与 TaskDec-PP 比高 0.82 /1.05。L4DR 是 2GPU、100ep、syncBN；PP /TaskDec 是 80ep，且训练精度和 warm start 不同。AP_R40 的 TaskDec best Moderate 来自 epoch73，不把其数值与官方 EAA 的 epoch79 当成同一个 checkpoint。',
            ['l4dr','vod'])
    table(s,.57,2.15,[3.5,3.5,2.6,2.6],['方法','来源 / checkpoint','EAA mAP','DC mAP'],[
        ['PP-Concat','本地 / epoch80','69.88','83.80'],
        ['TaskDec-PP','warm mild / epoch79','70.18','83.79'],
        ['L4DR','本地 / epoch99','71.00','84.84'],
        ['L4DR','文献 Table 3','72.70','87.47'],
    ],row_h=.64,font=19,highlight=[1])
    text(s,.75,5.68,11.75,.83,'训练差异：PP-Concat 用 AMP；TaskDec-PP 用 FP32 并 warm start；\nL4DR 本地复现为 2 GPU / 100 epochs / sync BN。',18,MUTED)

    s=slide('备份 E｜机制诊断的支持范围','BACKUP / 表征与控制读数',
            '中心距离是二维 PCA 诊断；前景 gate 统计与空间热力图属于不同证据。',
            '依据：168 帧分析导出、metrics.json；此处不推断传感器主导权切换',
            '左图是分别拟合的 PCA 空间中的模态中心距离，其数值不能作为严格的统计独立性证据。右图是采样 patch 的 gate 均值，不是所有 BEV patch 的全量均值。另一个重要事实是当前 reliability 偏 LiDAR，而 common/unique 的 Camera 前景绝对 cosine 约0.927，所以不能说所有分支已实现正交解耦。',
            ['pca','analysis_exports/taskdec_paper_visuals_260909/paper_visual_caption_notes.md'])
    picture(s,VIS_DIR/'paper_fig_taskdec_decoupling_and_gate.png',.58,2.0,12.17,3.82)
    for x,title,body in [(.58,'common 高维相似度','前景跨模态 cosine：0.948–0.962'),
                         (4.74,'不等于严格独立','Camera 的 |cos(c,u)| ≈ 0.927'),
                         (8.90,'reliability 偏 LiDAR','不支持“按天气切换主导模态”')]:
        box(s,x,6.02,3.84,.82,WHITE,LINE,True)
        text(s,x+.13,6.10,3.56,.31,title,16,TEAL,True)
        text(s,x+.13,6.50,3.56,.25,body,12,MUTED)

    s=slide('备份 F｜控制强度与辅助损失权重','BACKUP / 配置细节',
            '主模型 γ=0.75；关闭 reliability control 的移除结果来自单独选定的运行。',
            '依据：strength 比较、主配置；辅助损失外层 PATCH_DEC_WEIGHT=0.12',
            '非零强度0.5、0.75、1.0的 AP3D@0.3接近；0.0为组件移除，不能把表描述成同 checkpoint 的连续推理扫参。损失权重均还乘外层0.12：orth0.10、common0.12、unique0.015、gate0.15、task0.30。训练时 encoder 参数冻结，但 FREEZE_BN=False，状态更新设置必须明确。',
            ['strength','configs/ASF_task_dec_controlled_robust_v1_0.yml','configs/v1_0/cfg_A2F_scl_final.yml'])
    text(s,.62,2.13,5.8,.45,'控制强度 / 全量评测',22,INK,True)
    table(s,.57,2.81,[1.45,1.90,2.55],['γ','checkpoint','AP3D@0.3'],[
        ['0.0','model_9','79.62'],['0.5','model_0','88.19'],['0.75','model_0','88.36'],['1.0','model_0','88.33'],
    ],row_h=.57,font=18,highlight=[2])
    text(s,6.86,2.13,5.8,.45,'辅助损失 / 括号内权重',22,INK,True)
    table(s,6.81,2.81,[3.97,1.99],['项目','权重'],[
        ['common / unique 去相关','0.10'],['common 跨模态对齐','0.12'],['unique 分离','0.015'],
        ['foreground gate','0.15'],['task context','0.30'],
    ],row_h=.47,font=17)
    takeaway(s,'记录外层辅助权重 0.12，以及 FREEZE=True / FREEZE_BN=False。',size=17)

    assert len(prs.slides)==24
    # Geometry and provenance validation before saving.
    outside=[]
    for i,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            if sh.left < -1000 or sh.top < -1000 or sh.left+sh.width>prs.slide_width+1000 or sh.top+sh.height>prs.slide_height+1000:
                outside.append((i,sh.name))
        for ref in SLIDES[i-1]['sources']:
            assert (ROOT/ref).exists(),ref
    assert not outside,outside
    target=OUT/'TaskDec_现有结果与投稿讨论_260911.pptx'
    prs.save(target)
    metadata=dict(slides=len(prs.slides),main_slides=18,backup_slides=6,slide_size='16:9',
                  font=FONT,source_date='2026-09-11',weather_delta03=delta03,weather_delta05=delta05,
                  main_table=main_rows,ablation=ablation,availability=availability,pca=PCA_INFO,
                  sources=SOURCES,slides_with_notes=SLIDES,geometry_check='All slide objects within canvas',
                  generation='Existing local results only; no model training or inference')
    (OUT/'deck_manifest.json').write_text(json.dumps(metadata,indent=2,ensure_ascii=False)+'\n')
    md=['# TaskDec 汇报讲稿与来源','',
        '2026-09-11。前 18 页主讲，后 6 页备份；建议 15–20 分钟，按讨论节奏使用备份页。','']
    for record in SLIDES:
        md += [f"## {record['page']:02d}. {record['title']}",'',record['notes'],'','资料来源：','']
        md += [f'- [{Path(p).name}](../../{p})' for p in record['sources']]
        md += ['']
    (OUT/'汇报讲稿与来源.md').write_text('\n'.join(md)+'\n')
    print(f'Saved {target} ({target.stat().st_size/2**20:.2f} MiB); {len(prs.slides)} slides with speaker notes.')


if __name__ == '__main__':
    build()
